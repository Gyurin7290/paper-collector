#!/usr/bin/env python3
"""Search OpenAlex for papers in given journals/date range/keywords and download open-access PDFs.

Usage:
  python collect_papers.py search --field "peptide" --core "assembly,self-assembly,co-assembly" \
      --any "molecular dynamics,MD simulation,coarse-grained,conductivity,conductive,charge transport" \
      --venues "Nature Chemistry,ACS Nano,Angewandte Chemie International Edition,Journal of the American Chemical Society,Nature Communications,Science,Advanced Materials" \
      --from 2023-10-23 --to 2026-10-23 --out candidates.json [--fulltext] [--mailto you@uni.ac.kr]
  python collect_papers.py download --in curated.json --dest ./reference --mailto you@uni.ac.kr
  python collect_papers.py sheet --in curated.json --out paper_list.xlsx

search   -> candidates.json (list of dicts: doi,title,date,venue,oa,pdf,landing,abstract,match)
download -> PDFs named "YYYY_Journal_FirstWordsOfTitle.pdf" + updates 'status' in the json
sheet    -> Excel list (one row per paper, status column, DOI hyperlinks)
"""
import argparse, json, os, re, sys, time, urllib.parse, urllib.request

UA = "Mozilla/5.0 (paper-collector skill; mailto:{m})"
VENUE_ALIASES = {
    "angewandte chemie": "Angewandte Chemie International Edition",
    "angew": "Angewandte Chemie International Edition",
    "jacs": "Journal of the American Chemical Society",
    "nat commun": "Nature Communications",
    "nature communication": "Nature Communications",
    "nat chem": "Nature Chemistry",
    "adv mater": "Advanced Materials",
    "advanced material": "Advanced Materials",
}


def get(url, mailto, binary=False, timeout=60):
    req = urllib.request.Request(url, headers={"User-Agent": UA.format(m=mailto), "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        data = r.read()
        return (data, r.headers.get("Content-Type", "")) if binary else json.loads(data)


def resolve_sources(names, mailto):
    ids = {}
    for n in names:
        n = VENUE_ALIASES.get(n.strip().lower(), n.strip())
        r = get(f"https://api.openalex.org/sources?search={urllib.parse.quote(n)}&per_page=5"
                f"&select=id,display_name,works_count&mailto={mailto}", mailto)
        exact = [s for s in r["results"] if s["display_name"].lower() == n.lower()]
        best = (exact or sorted(r["results"], key=lambda s: -s["works_count"]))[:1]
        if best:
            ids[best[0]["display_name"]] = best[0]["id"].rsplit("/", 1)[1]
        else:
            print(f"[warn] venue not found: {n}", file=sys.stderr)
    return ids


def orq(terms):
    return "(" + " OR ".join(f'"{t}"' if " " in t or "-" in t else t for t in terms) + ")"


def abstract(w):
    inv = w.get("abstract_inverted_index") or {}
    pos = {}
    for word, idx in inv.items():
        for i in idx:
            pos[i] = word
    return " ".join(pos[i] for i in sorted(pos))


def search(a):
    venues = resolve_sources(a.venues.split(","), a.mailto)
    print("venues:", venues, file=sys.stderr)
    src = "|".join(venues.values())
    q = f"{a.field} AND {orq(a.core.split(','))} AND {orq(a.any.split(','))}"
    sel = "id,doi,title,publication_date,primary_location,open_access,best_oa_location,type,abstract_inverted_index"
    base = (f"https://api.openalex.org/works?filter=primary_location.source.id:{src},"
            f"from_publication_date:{a.from_},to_publication_date:{a.to},type:article")
    found = {}
    modes = [("title_abstract", f",title_and_abstract.search:{urllib.parse.quote(q)}")]
    if a.fulltext:
        modes.append(("fulltext", f"&search={urllib.parse.quote(q)}"))
    for tag, part in modes:
        for page in range(1, a.max_pages + 1):  # page= paging works without an API key (cursor needs one)
            u = f"{base}{part}&per_page=200&page={page}&select={sel}&mailto={a.mailto}"
            r = get(u, a.mailto)
            for w in r["results"]:
                d = (w.get("doi") or w["id"]).replace("https://doi.org/", "")
                if d in found:
                    continue
                bo = w.get("best_oa_location") or {}
                found[d] = dict(doi=d, title=w["title"], date=w["publication_date"],
                                venue=(w.get("primary_location") or {}).get("source", {}).get("display_name"),
                                oa=(w.get("open_access") or {}).get("oa_status"),
                                pdf=bo.get("pdf_url") or "", landing=bo.get("landing_page_url") or "",
                                abstract=abstract(w), match=tag)
            if len(r["results"]) < 200:
                break
            time.sleep(0.3)
    out = sorted(found.values(), key=lambda x: x["date"], reverse=True)
    json.dump(out, open(a.out, "w"), ensure_ascii=False, indent=1)
    print(f"{len(out)} candidates -> {a.out}")


def safe_name(p):
    short = {"Journal of the American Chemical Society": "JACS", "Angewandte Chemie International Edition": "Angew",
             "Nature Communications": "NatCommun", "Nature Chemistry": "NatChem", "Advanced Materials": "AdvMater",
             "ACS Nano": "ACSNano", "Science": "Science"}.get(p.get("venue"), re.sub(r"\W+", "", p.get("venue") or "X"))
    words = re.sub(r"[^\w\s-]", "", p["title"]).split()[:8]
    return f"{p['date'][:4]}_{short}_{'_'.join(words)}"[:120] + ".pdf"


def pdf_candidates(p, mailto):
    c = []
    if p.get("pdf"):
        c.append(p["pdf"])
    if p["doi"].startswith("10.1038/"):
        c.append(f"https://www.nature.com/articles/{p['doi'].split('/', 1)[1]}.pdf")
    if p["doi"].startswith("10.1002/"):
        c.append(f"https://onlinelibrary.wiley.com/doi/pdfdirect/{p['doi']}")
    if p["doi"].startswith("10.1021/"):
        c.append(f"https://pubs.acs.org/doi/pdf/{p['doi']}")
    try:  # Unpaywall: other legal OA copies (PMC, repositories, ChemRxiv)
        u = get(f"https://api.unpaywall.org/v2/{p['doi']}?email={mailto}", mailto, timeout=30)
        for loc in u.get("oa_locations") or []:
            if loc.get("url_for_pdf"):
                c.append(loc["url_for_pdf"])
    except Exception:
        pass
    seen, uniq = set(), []
    for x in c:
        if x not in seen:
            seen.add(x); uniq.append(x)
    return uniq


def download(a):
    papers = json.load(open(a.inp))
    os.makedirs(a.dest, exist_ok=True)
    for p in papers:
        if p.get("status") == "downloaded":
            continue
        if p.get("oa") == "closed" and not p.get("pdf"):
            p["status"] = "paywalled"
            continue
        p["status"] = "no_free_pdf"
        for url in pdf_candidates(p, a.mailto):
            try:
                data, ctype = get(url, a.mailto, binary=True)
                if data[:5] == b"%PDF-":  # publishers often return an HTML page instead of the PDF
                    fn = safe_name(p)
                    open(os.path.join(a.dest, fn), "wb").write(data)
                    p.update(status="downloaded", file=fn, pdf_used=url)
                    break
            except Exception as e:
                p.setdefault("errors", []).append(f"{url}: {e}"[:200])
            time.sleep(1)
        print(p["status"], p["title"][:70])
    json.dump(papers, open(a.inp, "w"), ensure_ascii=False, indent=1)


def sheet(a):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    papers = json.load(open(a.inp))
    label = {"downloaded": "다운로드 완료", "paywalled": "유료(구독 필요)", "no_free_pdf": "무료 PDF 받기 실패", "pending": "다운로드 대기"}
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "papers"
    ws.append(["No", "게재일", "저널", "제목", "키워드 매칭", "매칭 위치", "OA 상태", "파일명", "DOI", "상태"])
    for i, p in enumerate(papers, 1):
        ws.append([i, p["date"], p["venue"], p["title"], p.get("keyword", ""), p.get("match", ""), p.get("oa", ""),
                   p.get("file", ""), f"https://doi.org/{p['doi']}", label.get(p.get("status"), p.get("status", ""))])
        ws.cell(i + 1, 9).hyperlink = f"https://doi.org/{p['doi']}"
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for col, w in zip("ABCDEFGHIJ", [5, 11, 30, 80, 15, 14, 9, 45, 32, 18]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    wb.save(a.out); print("saved", a.out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    s = sp.add_parser("search")
    s.add_argument("--field", required=True); s.add_argument("--core", required=True); s.add_argument("--any", required=True)
    s.add_argument("--venues", required=True); s.add_argument("--from", dest="from_", required=True); s.add_argument("--to", required=True)
    s.add_argument("--out", default="candidates.json"); s.add_argument("--fulltext", action="store_true")
    s.add_argument("--max-pages", type=int, default=2); s.add_argument("--mailto", default="anonymous@example.org")
    d = sp.add_parser("download"); d.add_argument("--in", dest="inp", required=True); d.add_argument("--dest", required=True)
    d.add_argument("--mailto", default="anonymous@example.org")
    t = sp.add_parser("sheet"); t.add_argument("--in", dest="inp", required=True); t.add_argument("--out", required=True)
    a = ap.parse_args(); {"search": search, "download": download, "sheet": sheet}[a.cmd](a)
